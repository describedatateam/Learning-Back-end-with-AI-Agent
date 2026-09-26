"""Build study materials for each exercise: a slide deck (.pptx) and study notes (.html).

    python materials/build_materials.py            # all exercises
    python materials/build_materials.py 01 05      # only exercises whose folder starts with 01 or 05

Inputs
    materials/decks/<exercise>.json       slides written by the deck designer (format below)
    materials/resources/<exercise>.json   verified further-learning links
    learn/content/<exercise>/             the lesson (instructions.md) and meta.json
Outputs
    materials/out/decks/<exercise>.pptx
    materials/out/notes/<exercise>.html   (uploads to Google Drive as a Google Doc)

Deck JSON: {"exercise": "01-http-requests", "subtitle": "...", "slides": [...]}, where each
slide has a "type" and an optional "notes" (speaker notes):
    statement  title, text                       one big idea or analogy
    bullets    title, bullets[]                  3-5 short bullets
    terms      title, terms[{term, meaning}]     new words
    code       title, code, caption?, output?    one code example (max ~14 lines)
    compare    title, left{heading, bullets[]}, right{heading, bullets[]}
    steps      title, steps[{text, tests}]       the task, step by step
    scenario   title, text, bullets[]            a real-world use case
    check      title, questions[]                check-yourself questions
Text supports **bold** and `code`. The title slide and the "Keep learning" slide are added
automatically.
Requires: pip install -r materials/requirements.txt
"""
import html
import json
import re
import sys
from pathlib import Path

import markdown
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from pygments.lexers import PythonLexer
from pygments.token import Token

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / 'learn' / 'content'
DECKS = ROOT / 'materials' / 'decks'
RESOURCES = ROOT / 'materials' / 'resources'
OUT = ROOT / 'materials' / 'out'

INK, MUTED, LINE = RGBColor(0x1C, 0x23, 0x30), RGBColor(0x5D, 0x67, 0x78), RGBColor(0xDD, 0xE1, 0xE7)
PURPLE, BLUE, SOFT = RGBColor(0x8A, 0x4F, 0xE0), RGBColor(0x2F, 0x6F, 0xDF), RGBColor(0xF3, 0xEE, 0xFD)
DARK, CODE_BG, OUT_BG = RGBColor(0x11, 0x15, 0x1C), RGBColor(0x1E, 0x1E, 0x1E), RGBColor(0x2A, 0x2D, 0x33)
WHITE, INLINE_CODE = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x6A, 0x3F, 0xC0)
MONO, SANS = 'Consolas', 'Segoe UI'

# VS Code Dark+ colours for code.
CODE_COLOURS = [
    (Token.Comment, RGBColor(0x6A, 0x99, 0x55)),
    (Token.Keyword, RGBColor(0xC5, 0x86, 0xC0)),
    (Token.Name.Builtin, RGBColor(0x4E, 0xC9, 0xB0)),
    (Token.Name.Function, RGBColor(0xDC, 0xDC, 0xAA)),
    (Token.Name.Class, RGBColor(0x4E, 0xC9, 0xB0)),
    (Token.Name.Decorator, RGBColor(0xDC, 0xDC, 0xAA)),
    (Token.Literal.String, RGBColor(0xCE, 0x91, 0x78)),
    (Token.Literal.Number, RGBColor(0xB5, 0xCE, 0xA8)),
    (Token.Name, RGBColor(0x9C, 0xDC, 0xFE)),
    (Token.Operator.Word, RGBColor(0xC5, 0x86, 0xC0)),
]
PLAIN_CODE = RGBColor(0xD4, 0xD4, 0xD4)

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)


def code_colour(token_type):
    for kind, colour in CODE_COLOURS:
        if token_type in kind:
            return colour
    return PLAIN_CODE


def add_text(frame, text, size, colour=INK, bold=False, first=False, space_after=6, align=None):
    """Add a paragraph, turning **bold** and `code` into formatted runs."""
    para = frame.paragraphs[0] if first else frame.add_paragraph()
    para.space_after = Pt(space_after)
    if align:
        para.alignment = align
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        run = para.add_run()
        font = run.font
        font.size, font.color.rgb, font.bold, font.name = Pt(size), colour, bold, SANS
        if part.startswith('**'):
            run.text, font.bold = part[2:-2], True
        elif part.startswith('`'):
            run.text, font.name = part[1:-1], MONO
            if colour == INK:
                font.color.rgb = INLINE_CODE
        else:
            run.text = part
    return para


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return frame


def rounded(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def code_block(slide, code, left, top, width, height, label=None, background=CODE_BG):
    rounded(slide, left, top, width, height, background)
    lines = code.rstrip('\n').split('\n')
    size = 18 if len(lines) <= 8 else 16 if len(lines) <= 12 else 14 if len(lines) <= 16 else 12
    inset = Inches(0.3)
    frame = textbox(slide, left + inset, top + Inches(0.22), width - 2 * inset, height - Inches(0.4))
    if label:
        add_text(frame, label, 11, RGBColor(0x9A, 0xA4, 0xB5), bold=True, first=True, space_after=4)
    para = frame.add_paragraph() if label else frame.paragraphs[0]
    para.line_spacing = 1.1
    for token_type, value in PythonLexer().get_tokens(code.rstrip('\n')):
        pieces = value.split('\n')
        for i, piece in enumerate(pieces):
            if i:
                para = frame.add_paragraph()
                para.line_spacing = 1.1
            if piece:
                run = para.add_run()
                run.text = piece
                run.font.name, run.font.size = MONO, Pt(size)
                run.font.color.rgb = PLAIN_CODE if background != CODE_BG else code_colour(token_type)


class DeckBuilder:
    def __init__(self, exercise_dir, meta, deck):
        self.meta, self.deck, self.folder = meta, deck, exercise_dir.name
        self.number = int(self.folder.split('-')[0])
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    def new_slide(self, title=None, notes=None, dark=False):
        slide = self.prs.slides.add_slide(self.blank)
        self.count += 1
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = DARK if dark else WHITE
        if not dark:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = PURPLE
            bar.line.fill.background()
            footer = textbox(slide, MARGIN, H - Inches(0.5), W - 2 * MARGIN, Inches(0.3))
            add_text(footer, f'Backend Lab  ·  Exercise {self.number:02d}  ·  {self.meta["title"]}', 11, MUTED, first=True)
            num = textbox(slide, W - MARGIN - Inches(1), H - Inches(0.5), Inches(1), Inches(0.3))
            add_text(num, str(self.count), 11, MUTED, first=True, align=PP_ALIGN.RIGHT)
        if title:
            frame = textbox(slide, MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(0.9), MSO_ANCHOR.MIDDLE)
            add_text(frame, title, 32, WHITE if dark else INK, bold=True, first=True)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def body_frame(self, slide, top=Inches(1.55)):
        return textbox(slide, MARGIN, top, W - 2 * MARGIN, H - top - Inches(0.8))

    # ---- slide types ----
    def title_slide(self):
        slide = self.new_slide(dark=True, notes=self.deck.get('notes'))
        tag = textbox(slide, MARGIN, Inches(1.7), Inches(9), Inches(0.4))
        add_text(tag, f'WEEK {self.meta["week"]}  ·  EXERCISE {self.number:02d}  ·  {self.meta["concept"].upper()}',
                 14, RGBColor(0xB3, 0x8C, 0xFF), bold=True, first=True)
        frame = textbox(slide, MARGIN, Inches(2.3), W - 2 * MARGIN, Inches(2.2))
        add_text(frame, self.meta['title'], 54, WHITE, bold=True, first=True)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(4.55), Inches(1.4), Inches(0.08))
        line.fill.solid()
        line.fill.fore_color.rgb = PURPLE
        line.line.fill.background()
        sub = textbox(slide, MARGIN, Inches(4.9), W - 2 * MARGIN, Inches(1.2))
        add_text(sub, self.deck.get('subtitle') or self.meta['summary'], 22, RGBColor(0xC9, 0xD1, 0xDE), first=True)
        brand = textbox(slide, MARGIN, H - Inches(0.9), Inches(6), Inches(0.4))
        add_text(brand, 'Backend Lab', 14, RGBColor(0x9A, 0xA4, 0xB5), bold=True, first=True)

    def statement(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        rounded(slide, MARGIN, Inches(1.7), W - 2 * MARGIN, Inches(4.6), SOFT)
        frame = textbox(slide, MARGIN + Inches(0.5), Inches(1.9), W - 2 * MARGIN - Inches(1), Inches(4.2), MSO_ANCHOR.MIDDLE)
        for i, paragraph in enumerate(s['text'].split('\n\n')):
            add_text(frame, paragraph, 24, first=i == 0, space_after=14)

    def bullets(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        frame = self.body_frame(slide, Inches(1.75))
        for i, item in enumerate(s['bullets']):
            add_text(frame, f'•  {item}', 24, first=i == 0, space_after=16)

    def terms(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        rows = s['terms']
        row_h = min(Inches(1.0), int((H - Inches(2.5)) / max(len(rows), 1)))
        top = Inches(1.65)
        for i, row in enumerate(rows):
            y = top + i * row_h
            left = textbox(slide, MARGIN, y, Inches(3.3), row_h, MSO_ANCHOR.MIDDLE)
            add_text(left, row['term'], 22, PURPLE, bold=True, first=True)
            right = textbox(slide, MARGIN + Inches(3.5), y, W - 2 * MARGIN - Inches(3.5), row_h, MSO_ANCHOR.MIDDLE)
            add_text(right, row['meaning'], 19, first=True)
            if i:
                sep = slide.shapes.add_connector(1, MARGIN, y, W - MARGIN, y)
                sep.line.color.rgb = LINE

    def code(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        top = Inches(1.6)
        if s.get('caption'):
            frame = textbox(slide, MARGIN, top, W - 2 * MARGIN, Inches(0.9))
            add_text(frame, s['caption'], 20, first=True)
            top += Inches(0.95)
        bottom = H - Inches(0.75)
        output = s.get('output')
        out_lines = output.rstrip('\n').count('\n') + 1 if output else 0
        out_h = Inches(0.55 + 0.32 * out_lines) if output else 0
        gap = Inches(0.2) if output else 0
        code_block(slide, s['code'], MARGIN, top, W - 2 * MARGIN, bottom - top - out_h - gap)
        if output:
            code_block(slide, output, MARGIN, bottom - out_h, W - 2 * MARGIN, out_h, label='OUTPUT', background=OUT_BG)

    def compare(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        col_w = (W - 2 * MARGIN - Inches(0.4)) / 2
        for i, side in enumerate([s['left'], s['right']]):
            x = MARGIN + i * (col_w + Inches(0.4))
            rounded(slide, x, Inches(1.65), col_w, Inches(4.8), SOFT if i == 0 else RGBColor(0xE5, 0xEE, 0xFC))
            frame = textbox(slide, x + Inches(0.35), Inches(1.9), col_w - Inches(0.7), Inches(4.4))
            add_text(frame, side['heading'], 24, PURPLE if i == 0 else BLUE, bold=True, first=True, space_after=14)
            for item in side['bullets']:
                add_text(frame, f'•  {item}', 19, space_after=10)

    def steps(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        rows = s['steps']
        row_h = min(Inches(0.95), int((H - Inches(2.45)) / max(len(rows), 1)))
        for i, step in enumerate(rows):
            y = Inches(1.65) + i * row_h
            dot_top = y + int((row_h - Inches(0.5)) / 2)  # centre the number on its row, like the text
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN, dot_top, Inches(0.5), Inches(0.5))
            dot.fill.solid()
            dot.fill.fore_color.rgb = PURPLE
            dot.line.fill.background()
            dot.text_frame.text = str(i + 1)
            para = dot.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            para.runs[0].font.size, para.runs[0].font.bold, para.runs[0].font.color.rgb = Pt(16), True, WHITE
            frame = textbox(slide, MARGIN + Inches(0.75), y, W - 2 * MARGIN - Inches(2.9), row_h, MSO_ANCHOR.MIDDLE)
            add_text(frame, step['text'], 19, first=True)
            if step.get('tests'):
                tag = textbox(slide, W - MARGIN - Inches(2), y, Inches(2), row_h, MSO_ANCHOR.MIDDLE)
                add_text(tag, f'→ {step["tests"]}', 15, BLUE, bold=True, first=True, align=PP_ALIGN.RIGHT)

    def scenario(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        rounded(slide, MARGIN, Inches(1.65), W - 2 * MARGIN, Inches(1.7), RGBColor(0xE3, 0xF4, 0xEA))
        frame = textbox(slide, MARGIN + Inches(0.4), Inches(1.8), W - 2 * MARGIN - Inches(0.8), Inches(1.4), MSO_ANCHOR.MIDDLE)
        add_text(frame, s['text'], 21, first=True)
        frame = textbox(slide, MARGIN, Inches(3.65), W - 2 * MARGIN, Inches(3))
        for i, item in enumerate(s.get('bullets', [])):
            add_text(frame, f'•  {item}', 20, first=i == 0, space_after=12)

    def check(self, s):
        slide = self.new_slide(s['title'], s.get('notes'))
        frame = self.body_frame(slide, Inches(1.75))
        for i, question in enumerate(s['questions']):
            add_text(frame, f'{i + 1}.  {question}', 23, first=i == 0, space_after=18)

    def resources(self, items):
        slide = self.new_slide('Keep learning')
        frame = self.body_frame(slide, Inches(1.6))
        for i, item in enumerate(items[:5]):
            para = add_text(frame, '', 20, first=i == 0, space_after=2)
            run = para.add_run()
            run.text = item['title']
            run.font.size, run.font.bold, run.font.name, run.font.color.rgb = Pt(20), True, SANS, BLUE
            run.hyperlink.address = item['url']
            meta = f'{item.get("source", "")} · {item.get("kind", "")} · {item.get("level", "")}'.strip(' ·')
            add_text(frame, f'{meta}. {item.get("why", "")}', 15, MUTED, space_after=12)

    def build(self, resources):
        self.title_slide()
        for s in self.deck['slides']:
            getattr(self, s['type'])(s)
        if resources:
            self.resources(resources)
        return self.prs


NOTES_CSS = """
body { font-family: 'Segoe UI', Arial, sans-serif; line-height: 1.55; color: #1c2330; max-width: 820px; }
code, pre { font-family: Consolas, monospace; }
pre { background: #f3f4f6; padding: 10px 14px; border-radius: 6px; }
table { border-collapse: collapse; } th, td { border: 1px solid #dde1e7; padding: 4px 8px; text-align: left; }
"""


def build_notes(exercise_dir, meta, resources):
    lesson = (exercise_dir / 'instructions.md').read_text(encoding='utf-8')
    number = int(exercise_dir.name.split('-')[0])
    body = markdown.markdown(lesson, extensions=['fenced_code', 'tables', 'sane_lists'])
    links = ''.join(
        f'<li><a href="{html.escape(r["url"])}">{html.escape(r["title"])}</a> '
        f'({html.escape(r.get("source", ""))}, {html.escape(r.get("level", ""))}). {html.escape(r.get("why", ""))}</li>'
        for r in resources
    )
    return (
        f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(meta["title"])}</title>'
        f'<style>{NOTES_CSS}</style></head><body>'
        f'<p><strong>Backend Lab · Week {meta["week"]} · Exercise {number:02d} · {html.escape(meta["concept"])}</strong></p>'
        f'{body}'
        + (f'<h2>Keep learning</h2><ul>{links}</ul>' if links else '')
        + '</body></html>'
    )


def main(prefixes):
    (OUT / 'decks').mkdir(parents=True, exist_ok=True)
    (OUT / 'notes').mkdir(parents=True, exist_ok=True)
    for exercise_dir in sorted(p for p in CONTENT.iterdir() if p.is_dir()):
        if prefixes and not exercise_dir.name.startswith(tuple(prefixes)):
            continue
        meta = json.loads((exercise_dir / 'meta.json').read_text(encoding='utf-8'))
        resources_file = RESOURCES / f'{exercise_dir.name}.json'
        resources = json.loads(resources_file.read_text(encoding='utf-8'))['resources'] if resources_file.exists() else []

        notes_path = OUT / 'notes' / f'{exercise_dir.name}.html'
        notes_path.write_text(build_notes(exercise_dir, meta, resources), encoding='utf-8')
        deck_file = DECKS / f'{exercise_dir.name}.json'
        if deck_file.exists():
            deck = json.loads(deck_file.read_text(encoding='utf-8'))
            deck_path = OUT / 'decks' / f'{exercise_dir.name}.pptx'
            DeckBuilder(exercise_dir, meta, deck).build(resources).save(deck_path)
            print(f'{exercise_dir.name}: deck ({len(deck["slides"]) + 1 + bool(resources)} slides), notes')
        else:
            print(f'{exercise_dir.name}: notes only (no deck spec yet)')


if __name__ == '__main__':
    main(sys.argv[1:])
